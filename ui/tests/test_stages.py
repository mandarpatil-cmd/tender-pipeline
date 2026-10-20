"""Stage pages, their exports, and a company that is already known."""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pipeline_core.db import engine, ensure_schema, session
from pipeline_core.grid import award_view
from pipeline_core.models import Vendor
from pipeline_core.queries import (
    mark_enriched,
    record_outreach,
    replace_awards,
    record_pdf_contacts,
    upsert_tender,
    upsert_vendor,
    utcnow,
)
from starlette.datastructures import QueryParams

from ops_ui.app import create_app
from ops_ui.present import STAGE_COLUMNS, stage_query


def test_stage_export_matches_the_query_including_later_pages(tmp_path, monkeypatch):
    path = tmp_path / "pipeline.sqlite3"
    bind = engine(path)
    ensure_schema(bind)
    with session(bind) as current:
        for index in range(3):
            vendor_id = upsert_vendor(current, name_raw=f"Company {index}")
            tender_id = f"T-{index}"
            upsert_tender(current, tender_id=tender_id, title=tender_id, scraped_at=utcnow())
            replace_awards(
                current,
                tender_id,
                [{"bid_number": "1", "vendor_id": vendor_id, "bidder_name": f"Company {index}"}],
            )
    monkeypatch.setattr("ops_ui.routes.awards.PAGE_SIZE", 1)
    client = TestClient(create_app(path))
    page = client.get("/")
    assert page.text.count('class="open"') == 1

    workbook = load_workbook(BytesIO(client.get("/export.xlsx?page=2").content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == [
        "Tender",
        "Title",
        "Organisation",
        "Status",
        "Contract date",
        "Contract value",
        "Contract currency",
        "Bid",
        "Rank",
        "Quoted",
        "Awarded",
        "Currency",
        "Work title",
        "Name",
        "City",
        "State",
        "Email",
        "Phone",
        "PDF email",
        "PDF phone",
        "Source",
        "Scraped",
        "Origin",
        "Portal search",
        "Page",
        "S.No",
        "Run",
    ]
    assert sheet.max_row == 4
    query = stage_query("scrape", QueryParams("page=2"))
    with session(engine(path)) as current:
        view = award_view(current, query, page_size=None)
    assert len(view.rows) == 3
    assert [sheet.cell(row, 1).value for row in range(2, 5)] == [row["tender_id"] for row in view.rows]


def test_a_later_tender_reuses_the_stored_contact(tmp_path):
    path = tmp_path / "pipeline.sqlite3"
    bind = engine(path)
    ensure_schema(bind)
    with session(bind) as current:
        known = upsert_vendor(current, name_raw="Known Roads Ltd")
        waiting = upsert_vendor(current, name_raw="Still Waiting Ltd")
        mark_enriched(current, known, email="known@example.com", phone=None)
        upsert_tender(current, tender_id="T-1", title="First win", scraped_at=utcnow())
        upsert_tender(current, tender_id="T-2", title="Second win", scraped_at=utcnow())
        upsert_tender(current, tender_id="T-3", title="Waiting win", scraped_at=utcnow())
        replace_awards(
            current,
            "T-1",
            [{"bid_number": "1", "vendor_id": known, "bidder_name": "Known Roads Ltd"}],
        )
        replace_awards(
            current,
            "T-2",
            [{"bid_number": "1", "vendor_id": known, "bidder_name": "Known Roads Ltd"}],
        )
        replace_awards(
            current,
            "T-3",
            [{"bid_number": "1", "vendor_id": waiting, "bidder_name": "Still Waiting Ltd"}],
        )
        record_outreach(
            current,
            vendor_id=known,
            tender_id="T-1",
            email="known@example.com",
            subject="sent",
            transport="gmail",
            status="sent",
        )
    client = TestClient(create_app(path))
    scrape = client.get("/").text
    assert ">T-2</a>" in scrape
    assert "known@example.com" in scrape
    enrich = client.get("/enrich").text
    assert "Still Waiting Ltd" in enrich
    assert "Known Roads" not in enrich
    assert ">T-2</a>" not in enrich
    mail = client.get("/mail").text
    assert ">T-2</a>" in mail
    assert "known@example.com" in mail
    assert ">T-1</a>" not in mail
    confirm = client.post("/selection", data={"action": "lookup", "vendor_id": f"{known}:T-2"})
    assert "Skipped, because they already have a contact" in confirm.text
    assert "Known Roads Ltd" in confirm.text
    assert f'value="{known}"' not in confirm.text
    with session(engine(path)) as current:
        assert current.get(Vendor, known).email == "known@example.com"
        assert current.get(Vendor, waiting).enrichment_status == "pending"


def test_phone_only_company_is_absent_from_mail(tmp_path):
    path = tmp_path / "pipeline.sqlite3"
    ensure_schema(engine(path))
    with session(engine(path)) as current:
        phoned = upsert_vendor(current, name_raw="Phone Only Traders")
        mark_enriched(current, phoned, email=None, phone="+919876543210")
        upsert_tender(current, tender_id="T-9", title="Call", scraped_at=utcnow())
        replace_awards(
            current,
            "T-9",
            [{"bid_number": "1", "vendor_id": phoned, "bidder_name": "Phone Only Traders"}],
        )
    html = TestClient(create_app(path)).get("/mail").text
    assert "Phone Only Traders" not in html
    assert "No awards match." in html


def test_pdf_contact_keeps_only_companies_with_both_clues(tmp_path):
    path = tmp_path / "pipeline.sqlite3"
    bind = engine(path)
    ensure_schema(bind)
    with session(bind) as current:
        both = upsert_vendor(current, name_raw="Both Clues Ltd")
        email_only = upsert_vendor(current, name_raw="Email Clue Ltd")
        neither = upsert_vendor(current, name_raw="No Clue Ltd")
        for tender_id, vendor_id, name in (
            ("BOTH", both, "Both Clues Ltd"),
            ("ONE", email_only, "Email Clue Ltd"),
            ("NONE", neither, "No Clue Ltd"),
        ):
            upsert_tender(current, tender_id=tender_id, title=name, scraped_at=utcnow())
            replace_awards(
                current,
                tender_id,
                [{"bid_number": "1", "vendor_id": vendor_id, "bidder_name": name}],
            )
        # PDF clues belong to the tender whose files they were read from.
        record_pdf_contacts(current, both, "BOTH", ["both@example.com"], ["9123456780"])
        record_pdf_contacts(current, email_only, "ONE", ["one@example.com"], [])
    client = TestClient(create_app(path))

    scrape = client.get("/?pdf_contact=yes")
    assert "Both Clues Ltd" in scrape.text
    assert "Email Clue Ltd" not in scrape.text
    assert "No Clue Ltd" not in scrape.text
    assert "PDF email" in scrape.text

    enrich = client.get("/enrich?pdf_contact=yes")
    assert "Both Clues Ltd" in enrich.text
    assert "Email Clue Ltd" not in enrich.text
    assert "PDF email" in enrich.text
    assert "PDF phone" in enrich.text

    mail = client.get("/mail")
    assert 'name="pdf_contact"' not in mail.text


def test_scrape_table_shows_where_each_tender_came_from(tmp_path):
    path = tmp_path / "pipeline.sqlite3"
    ensure_schema(engine(path))
    with session(engine(path)) as current:
        vendor_id = upsert_vendor(current, name_raw="Portal Roads Ltd")
        upsert_tender(
            current,
            tender_id="T-PORTAL",
            title="Road",
            scraped_at=utcnow(),
            origin="eprocure.gov.in",
            search_mode="awarded",
            listing_page=2,
            listing_serial=14,
            scrape_job_id=7,
        )
        replace_awards(
            current,
            "T-PORTAL",
            [{"bid_number": "1", "vendor_id": vendor_id, "bidder_name": "Portal Roads Ltd"}],
        )
    html = TestClient(create_app(path)).get("/").text

    for label in ("Origin", "Portal search", "Page", "S.No", "Run"):
        assert f">{label}</a>" in html
    assert 'data-col="origin">eprocure.gov.in</td>' in html
    assert 'data-col="search_mode">awarded</td>' in html
    assert 'data-col="listing_serial">14</td>' in html
    assert '<a href="/jobs/7">7</a>' in html


def test_enrich_table_shows_where_the_address_was_found(tmp_path):
    from pipeline_core.queries import mark_enriched, record_llm_run

    path = tmp_path / "pipeline.sqlite3"
    ensure_schema(engine(path))
    with session(engine(path)) as current:
        vendor_id = upsert_vendor(current, name_raw="Traced Roads")
        upsert_tender(current, tender_id="T-1", title="Cables", scraped_at=utcnow())
        replace_awards(
            current,
            "T-1",
            [{"bid_number": "1", "vendor_id": vendor_id, "bidder_name": "Traced Roads"}],
        )
        run_id = record_llm_run(
            current,
            vendor_id=vendor_id,
            model="test",
            prompt_version="contact-v1",
            input_json="{}",
            output_json='{"email": "found@co.example"}',
            status="ok",
            job_id=9,
        )
        record_llm_run(
            current,
            vendor_id=vendor_id,
            model="test",
            prompt_version="contact-v1",
            input_json="{}",
            output_json=None,
            status="error",
            error="402",
        )
        mark_enriched(
            current,
            vendor_id,
            email="found@co.example",
            phone=None,
            source_url="https://co.example/about",
            llm_run_id=run_id,
        )
    html = TestClient(create_app(path)).get("/enrich?enrichment_status=done").text

    assert ">Found at</a>" in html
    assert ">Lookups</a>" in html
    assert 'data-col="found_at"' in html
    assert "https://co.example/about" in html
    assert 'data-col="lookups">2</td>' in html


def test_stage_column_sets_are_the_ones_the_pages_export():
    for key in ("origin", "search_mode", "listing_page", "listing_serial", "scrape_job_id"):
        assert key in STAGE_COLUMNS["scrape"]
        assert key not in STAGE_COLUMNS["enrich"]
    assert "found_at" in STAGE_COLUMNS["enrich"]
    assert "lookups" in STAGE_COLUMNS["enrich"]
    assert "found_at" not in STAGE_COLUMNS["scrape"]
    assert "email" in STAGE_COLUMNS["scrape"]
    assert "pdf_email" in STAGE_COLUMNS["scrape"]
    assert "pdf_phone" in STAGE_COLUMNS["scrape"]
    assert "contract_value" in STAGE_COLUMNS["scrape"]
    assert "awarded_value" in STAGE_COLUMNS["scrape"]
    assert "pdf_email" in STAGE_COLUMNS["enrich"]
    assert "organisation" in STAGE_COLUMNS["enrich"]
    assert "contract_date" in STAGE_COLUMNS["enrich"]
    assert "contract_value" in STAGE_COLUMNS["enrich"]
    assert "source" in STAGE_COLUMNS["enrich"]
    assert "gstin" not in STAGE_COLUMNS["scrape"]
    assert "outreach_status" in STAGE_COLUMNS["mail"]
    assert "organisation" in STAGE_COLUMNS["mail"]
    assert "contract_date" in STAGE_COLUMNS["mail"]
    assert "contract_value" in STAGE_COLUMNS["mail"]
    assert "scraped_at" in STAGE_COLUMNS["mail"]
