"""Stage exports and the runs list follow the typed columns."""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from ops_ui.app import create_app
from ops_ui.jobs import JobQuery, list_jobs
from pipeline_core.db import engine, ensure_schema, session
from pipeline_core.grid import AwardQuery, award_view
from pipeline_core.models import UiJob
from pipeline_core.queries import replace_awards, upsert_tender, upsert_vendor, utcnow


def _seed(path):
    bind = engine(path)
    ensure_schema(bind)
    with session(bind) as current:
        low = upsert_vendor(current, name_raw="Low Bid Ltd")
        high = upsert_vendor(current, name_raw="High Bid Ltd")
        upsert_tender(
            current,
            tender_id="LOW",
            title="Small",
            contract_date="01-Jul-2026",
            contract_value="INR 179,853.24",
            scraped_at=utcnow(),
        )
        upsert_tender(
            current,
            tender_id="HIGH",
            title="Large",
            contract_date="21-Sep-2026",
            contract_value="INR 4,608,100",
            scraped_at=utcnow(),
        )
        upsert_tender(
            current,
            tender_id="BIG",
            title="Biggest",
            contract_value="10000000",
            scraped_at=utcnow(),
        )
        replace_awards(
            current,
            "LOW",
            [{"bid_number": "1", "vendor_id": low, "bidder_name": "Low Bid Ltd"}],
        )
        replace_awards(
            current,
            "HIGH",
            [
                {
                    "bid_number": "1",
                    "vendor_id": high,
                    "bidder_name": "High Bid Ltd",
                    "quoted_value": "46,08,100.00",
                }
            ],
        )
        replace_awards(
            current,
            "BIG",
            [{"bid_number": "1", "vendor_id": high, "bidder_name": "High Bid Ltd"}],
        )


def test_scrape_excel_follows_the_numeric_contract_order(tmp_path):
    path = tmp_path / "pipeline.sqlite3"
    _seed(path)
    client = TestClient(create_app(path))
    sheet = load_workbook(
        BytesIO(client.get("/scrape/export.xlsx?sort=contract_value&dir=asc").content)
    ).active
    headers = [cell.value for cell in sheet[1]]
    ids = [sheet.cell(row, 1).value for row in range(2, 5)]

    with session(engine(path)) as current:
        view = award_view(
            current,
            AwardQuery(sort="contract_value", direction="asc"),
            page_size=None,
        )
    assert ids == [row["tender_id"] for row in view.rows]
    assert ids == ["LOW", "HIGH", "BIG"]
    value_at = headers.index("Contract value") + 1
    currency_at = headers.index("Contract currency") + 1
    assert sheet.cell(2, value_at).value == "179,853.24"
    assert sheet.cell(2, currency_at).value == "INR"
    assert sheet.cell(3, value_at).value == "4,608,100.00"


def test_runs_follow_stage_state_and_the_started_day(tmp_path):
    path = tmp_path / "pipeline.sqlite3"
    bind = engine(path)
    ensure_schema(bind)
    with session(bind) as current:
        current.add(
            UiJob(
                stage="scrape",
                state="done",
                operator="local",
                log="",
                started_at="2026-07-01T10:00:00+00:00",
            )
        )
        current.add(
            UiJob(
                stage="enrich",
                state="failed",
                operator="local",
                log="",
                started_at="2026-09-21T10:00:00+00:00",
            )
        )
        current.add(
            UiJob(
                stage="outreach",
                state="done",
                operator="local",
                log="",
                started_at="2026-09-22T10:00:00+00:00",
            )
        )

    assert [row.stage for row in list_jobs(path, JobQuery(stage="enrich")).rows] == ["enrich"]
    assert [row.state for row in list_jobs(path, JobQuery(state="failed")).rows] == ["failed"]
    september = list_jobs(
        path, JobQuery(started_from="2026-09-01", started_to="2026-09-21")
    )
    assert [row.stage for row in september.rows] == ["enrich"]
    later = list_jobs(path, JobQuery(started_from="2026-09-22"))
    assert [row.stage for row in later.rows] == ["outreach"]
