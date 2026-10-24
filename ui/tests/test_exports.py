"""Excel and CSV hold the rows and columns on screen, and change nothing (FR-105)."""

from __future__ import annotations

import csv
import io
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pipeline_core.db import engine, ensure_schema, session, table_counts
from pipeline_core.models import UiJob
from pipeline_core.queries import mark_enriched, replace_awards, upsert_tender, upsert_vendor, utcnow

from ops_ui.app import create_app
from ops_ui.export import layout
from ops_ui.present import STAGE_COLUMNS

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _seed(path) -> int:
    """A mailable company from one scrape run, and two from an older run."""
    bind = engine(path)
    ensure_schema(bind)
    with session(bind) as current:
        job = UiJob(stage="scrape", state="done", operator="local", params_json="{}", log="",
                    started_at=utcnow(), finished_at=utcnow())
        current.add(job)
        current.flush()
        companies = [
            ("T-MAIL", "Mail Works", "a@mail.example", None, job.job_id),
            ("T-WAIT", "Pending Works", None, None, 9),
            ("T-CALL", "Phone Only Works", None, "9876543210", 9),
        ]
        for tender_id, name, email, phone, run in companies:
            vendor_id = upsert_vendor(current, name_raw=name, state="Goa")
            if email or phone:
                mark_enriched(current, vendor_id, email=email, phone=phone)
            upsert_tender(current, tender_id=tender_id, title=f"Work {tender_id}",
                          scraped_at=utcnow(), scrape_job_id=run)
            replace_awards(current, tender_id,
                           [{"bid_number": "1", "vendor_id": vendor_id, "bidder_name": name}])
        return int(job.job_id)


@pytest.fixture
def seeded(tmp_path):
    path = tmp_path / "pipeline.sqlite3"
    job_id = _seed(path)
    return TestClient(create_app(path)), path, job_id


def _sheet(answer) -> list[list]:
    sheet = load_workbook(BytesIO(answer.content)).active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def _csv(answer) -> list[list[str]]:
    return list(csv.reader(io.StringIO(answer.content.decode("utf-8-sig"))))


@pytest.mark.parametrize("stage", ["enrich", "mail"])
def test_a_stage_export_holds_the_rows_and_columns_on_its_page(seeded, stage):
    client, _path, _job = seeded
    on_screen = client.get(f"/{stage}").text.count('class="open"')

    workbook = _sheet(client.get(f"/{stage}/export.xlsx"))
    sheet = _csv(client.get(f"/{stage}/export.csv"))

    assert on_screen > 0
    assert workbook[0] == layout(STAGE_COLUMNS[stage])[0] == sheet[0]
    assert len(workbook) - 1 == len(sheet) - 1 == on_screen


def test_a_filtered_stage_export_follows_the_filter(seeded):
    client, _path, _job = seeded
    page = client.get("/enrich?enrichment_status=done")
    workbook = _sheet(client.get("/enrich/export.xlsx?enrichment_status=done"))

    names = sorted(row[workbook[0].index("Name")] for row in workbook[1:])
    assert names == ["Mail Works", "Phone Only Works"]
    assert len(names) == page.text.count('class="open"')


def test_a_run_export_holds_only_that_runs_rows(seeded):
    client, _path, job_id = seeded
    page = client.get(f"/scrape?job={job_id}")
    workbook = _sheet(client.get(f"/scrape/export.xlsx?job={job_id}"))

    assert workbook[0] == layout(STAGE_COLUMNS["scrape"])[0]
    assert [row[workbook[0].index("Tender")] for row in workbook[1:]] == ["T-MAIL"]
    assert page.text.count('class="open"') == 1


def test_the_mail_export_names_the_address_that_would_be_mailed(seeded):
    client, _path, _job = seeded
    workbook = _sheet(client.get("/mail/export.xlsx"))
    header = workbook[0]
    by_tender = {row[header.index("Tender")]: row for row in workbook[1:]}

    assert "a@mail.example" in str(by_tender["T-MAIL"])
    assert "T-CALL" not in by_tender, "a phone alone is never mailed"


@pytest.mark.parametrize(
    ("url", "filename", "media"),
    [
        ("/export.xlsx", "awards.xlsx", _XLSX),
        ("/export.csv", "awards.csv", "text/csv"),
        ("/enrich/export.xlsx", "enrich.xlsx", _XLSX),
        ("/mail/export.csv", "mail.csv", "text/csv"),
    ],
)
def test_an_export_is_a_named_download(seeded, url, filename, media):
    client, _path, _job = seeded
    answer = client.get(url)

    assert answer.status_code == 200
    assert answer.headers["content-type"].startswith(media)
    assert answer.headers["content-disposition"] == f'attachment; filename="{filename}"'


@pytest.mark.parametrize("url", ["/home/export.xlsx", "/nope/export.csv", "/jobs/export.xlsx"])
def test_there_is_no_export_for_other_pages(seeded, url):
    client, _path, _job = seeded
    answer = client.get(url)

    assert answer.status_code in {404, 405}


def test_an_export_without_a_database_is_a_404(tmp_path):
    client = TestClient(create_app(tmp_path / "missing.sqlite3"))

    for url in ["/export.xlsx", "/export.csv", "/enrich/export.xlsx", "/mail/export.csv"]:
        assert client.get(url).status_code == 404, url
    assert not (tmp_path / "missing.sqlite3").exists()


def test_an_export_changes_nothing(seeded):
    client, path, job_id = seeded
    client.get("/")
    before = table_counts(engine(path))

    for url in ["/export.xlsx", "/export.csv", "/enrich/export.xlsx", "/mail/export.xlsx",
                f"/scrape/export.csv?job={job_id}"]:
        assert client.get(url).status_code == 200

    assert table_counts(engine(path)) == before
