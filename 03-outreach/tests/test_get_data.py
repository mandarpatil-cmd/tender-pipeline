"""The download must agree with what stage 3 would actually mail."""

from __future__ import annotations

import csv

import pytest
from pipeline_core.db import session
from pipeline_core.models import OUTREACH_FAILED, OUTREACH_SENT
from pipeline_core.queries import mark_enriched, record_outreach, upsert_vendor

import campaign
import get_data

pytest.importorskip("openpyxl")


def _vendor(name: str, email: str | None, *, source: str = "scrape") -> int:
    with session() as current:
        vendor_id = upsert_vendor(current, name_raw=name, source=source)
        mark_enriched(current, vendor_id, email=email, phone=None)
        return vendor_id


def test_the_export_is_the_same_list_stage_3_would_send():
    """The point of the whole script: review the real queue, not a lookalike."""
    _vendor("Kanta Enterprises", "hi@kanta.example")
    _vendor("Reachable Ltd", "hi@reachable.example")
    _vendor("No Address Ltd", None)

    exported = [r["vendor_id"] for r in get_data.collect_rows()]
    would_send = [t.vendor_id for t in campaign.load_recipients()]

    assert exported == would_send


def test_a_sent_vendor_drops_out_of_the_export_too():
    vendor_id = _vendor("Kanta Enterprises", "hi@kanta.example")
    _vendor("Still Waiting Ltd", "hi@waiting.example")

    with session() as current:
        record_outreach(
            current,
            vendor_id=vendor_id,
            email="hi@kanta.example",
            subject="Enquiry",
            transport="gmail",
            status=OUTREACH_SENT,
        )

    assert [r["company"] for r in get_data.collect_rows()] == ["Still Waiting Ltd"]


def test_include_sent_shows_the_whole_campaign_with_its_status():
    sent = _vendor("Mailed Ltd", "hi@mailed.example")
    _vendor("Waiting Ltd", "hi@waiting.example")

    with session() as current:
        record_outreach(
            current,
            vendor_id=sent,
            email="hi@mailed.example",
            subject="Enquiry",
            transport="gmail",
            status=OUTREACH_SENT,
        )

    rows = {r["company"]: r for r in get_data.collect_rows(include_sent=True)}
    assert set(rows) == {"Mailed Ltd", "Waiting Ltd"}
    assert rows["Mailed Ltd"]["outreach_status"] == OUTREACH_SENT
    assert rows["Mailed Ltd"]["last_attempt_at"]
    assert rows["Waiting Ltd"]["outreach_status"] == "queued"


def test_a_failed_attempt_still_counts_as_queued():
    """A failure is not a send, so that vendor is still going to be mailed."""
    vendor_id = _vendor("Bounces Ltd", "nope@bounces.example")
    with session() as current:
        record_outreach(
            current,
            vendor_id=vendor_id,
            email="nope@bounces.example",
            subject="Enquiry",
            transport="gmail",
            status=OUTREACH_FAILED,
            error="550",
        )

    row = get_data.collect_rows()[0]
    assert row["outreach_status"] == "queued"
    assert row["last_attempt_at"]  # but the attempt is still visible


def test_source_filter():
    _vendor("Scraped Ltd", "a@scraped.example", source="scrape")
    _vendor("Imported Ltd", "b@imported.example", source="bideasy")

    rows = get_data.collect_rows(source="scrape")
    assert [r["company"] for r in rows] == ["Scraped Ltd"]


def test_writes_both_files_with_every_column(tmp_path):
    _vendor("Kanta Enterprises", "hi@kanta.example")
    rows = get_data.collect_rows()

    csv_path = get_data.write_csv(rows, tmp_path / "queue.csv")
    xlsx_path = get_data.write_xlsx(rows, tmp_path / "queue.xlsx")

    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        read_back = list(csv.DictReader(handle))
    assert list(read_back[0]) == get_data.COLUMNS
    assert read_back[0]["email"] == "hi@kanta.example"

    from openpyxl import load_workbook

    sheet = load_workbook(xlsx_path).active
    assert [c.value for c in sheet[1]] == get_data.COLUMNS
    assert sheet.cell(row=2, column=get_data.COLUMNS.index("email") + 1).value == (
        "hi@kanta.example"
    )


def test_an_empty_queue_writes_a_header_only_file(tmp_path):
    csv_path = get_data.write_csv([], tmp_path / "queue.csv")
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        assert list(csv.reader(handle)) == [get_data.COLUMNS]
